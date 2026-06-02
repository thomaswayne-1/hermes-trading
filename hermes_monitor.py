#!/usr/bin/env python3
"""
hermes_monitor.py — combined live monitor + CSV tracker for Hermes Trading.

Replaces both watch_engine.py and excel_tracker.py.

Run:
    python3 hermes_monitor.py

Every 10 seconds:
  - Renders full engine monitor to terminal (regime, C/K, sub-signals,
    open positions, performance)
  - Writes position.csv, stats.csv, model.csv, trades.csv
  - Writes hermes_trades.xlsx  (Trades + Performance sheets, Arial font)

Environment (optional — auto-detected):
    RAILWAY_URL   default: https://hermes-trading-production-169f.up.railway.app
    API_SECRET    default: hermes2024
"""

import csv
import json
import os
import ssl
import time
import traceback
import urllib.request
from datetime import datetime, timezone
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.styles import (
    Alignment, Border, Font, PatternFill, Side
)
from openpyxl.utils import get_column_letter

# ── Config ────────────────────────────────────────────────────────────────────
RAILWAY_URL = os.getenv("RAILWAY_URL", "https://hermes-trading-production-169f.up.railway.app").rstrip("/")
API_SECRET  = os.getenv("API_SECRET",  "hermes2024")
STARTING    = 100_000.0
INTERVAL    = 10
BASE        = Path(__file__).parent
W           = 60   # terminal display width

_SSL = ssl.create_default_context()
_SSL.check_hostname = False
_SSL.verify_mode    = ssl.CERT_NONE


# ── HTTP ──────────────────────────────────────────────────────────────────────

def _get(path: str):
    req = urllib.request.Request(
        f"{RAILWAY_URL}{path}",
        headers={"X-API-Key": API_SECRET},
    )
    with urllib.request.urlopen(req, context=_SSL, timeout=10) as r:
        return json.loads(r.read())


def _live_btc_price() -> float:
    """
    Fetch a live BTC/USD price directly from Kraken's public ticker.
    Used when the Railway heartbeat is stale (server-side loop down).
    Returns 0.0 on any error.
    """
    try:
        req = urllib.request.Request(
            "https://api.kraken.com/0/public/Ticker?pair=XBTUSD",
            headers={"User-Agent": "hermes-monitor/1.0"},
        )
        with urllib.request.urlopen(req, context=_SSL, timeout=5) as r:
            data = json.loads(r.read())
        result = data.get("result", {})
        key = next((k for k in result if k != "last"), None)
        if key:
            return float(result[key]["c"][0])
    except Exception:
        pass
    return 0.0


# ── Last-known-good cache ─────────────────────────────────────────────────────

_cache: dict = {
    "state": {}, "trades": [],
    "stale": False, "stale_since": None,
}


def fetch_all():
    global _cache
    try:
        state  = _get("/state")
        trades = _get("/trades")
        if not isinstance(trades, list):
            trades = trades.get("trades", [])
        closed = [t for t in trades if t.get("closed")]

        # Never trust a response with fewer closed trades than cache
        if len(closed) < len(_cache["trades"]):
            raise ValueError(
                f"Response has {len(closed)} closed trades but cache has "
                f"{len(_cache['trades'])} — likely a redeploy blip"
            )

        _cache.update({"state": state, "trades": closed,
                       "stale": False, "stale_since": None})
    except Exception as e:
        now = datetime.now().strftime("%H:%M:%S")
        if not _cache["stale"]:
            _cache["stale"] = True
            _cache["stale_since"] = now
        print(f"  [{now}] ⚠ Using cached data — {e}")

    return _cache["state"], _cache["trades"]


# ── Helpers ───────────────────────────────────────────────────────────────────

def bar(value: float, width: int = 22) -> str:
    mid    = width // 2
    filled = min(int(abs(value) * mid + 0.5), mid)
    if value >= 0:
        l, r = " " * mid, "█" * filled + " " * (mid - filled)
    else:
        l, r = " " * (mid - filled) + "█" * filled, " " * mid
    return f"[{l}│{r}]"


def _net(t: dict) -> float:
    v = t.get("pnl_pct_net")
    return v if v is not None else t.get("pnl_pct_levered", 0)


def _calc_balance(closed: list) -> tuple[float, list]:
    """Cumulative balance + per-trade dollar PnL list. Identical logic in both tools."""
    bal   = STARTING
    dolls = []
    for t in closed:
        size = float(t.get("position_size_r", 0.15))
        d    = _net(t) * size * bal
        dolls.append(d)
        bal  += d
    return bal, dolls


def _usd(v: float) -> str:
    return f"{'+' if v >= 0 else '-'}${abs(v):,.2f}"


def _pct(v: float) -> str:
    return f"{'+' if v > 0 else ''}{v * 100:.2f}%"


def _age(ts: str) -> str:
    try:
        t = datetime.fromisoformat(ts)
        if t.tzinfo is None:
            t = t.replace(tzinfo=timezone.utc)
        s = int((datetime.now(timezone.utc) - t).total_seconds())
        if s < 60:   return f"{s}s ago"
        if s < 3600: return f"{s // 60}m ago"
        return f"{s // 3600}h {(s % 3600) // 60}m ago"
    except Exception:
        return "?"


def _write_csv(name: str, rows: list) -> None:
    with open(BASE / name, "w", newline="") as f:
        csv.writer(f).writerows(rows)


# ── Terminal render ───────────────────────────────────────────────────────────

TW = 100   # total terminal width

def _row(label: str, value: str, w: int = TW) -> str:
    """Left-label, right-value row padded to width."""
    gap = w - 4 - len(label) - len(value)
    return f"  {label}{' ' * max(1, gap)}{value}"


def render_terminal(state: dict, closed: list) -> None:
    hb    = state.get("heartbeat", {})
    eng   = hb.get("engine", {})
    cb    = hb.get("circuit_breaker", {})
    strat = state.get("strategy", {})
    sigs  = eng.get("sub_signals") or {}
    wts   = eng.get("weights_used") or {}

    price  = float(hb.get("price",        0))
    rsi    = float(hb.get("rsi",          50))
    macd   = float(hb.get("macd_hist",    0))
    bb     = float(hb.get("bb_pct",       0.5))
    atr    = float(hb.get("atr",          0))
    ob_imb = float(hb.get("ob_imbalance", 0))

    # If the server heartbeat is stale, fetch a live BTC price directly
    # from Kraken so the monitor always shows an accurate price.
    _ts_raw_pre = hb.get("ts", "")
    try:
        _dt_pre = datetime.fromisoformat(_ts_raw_pre)
        if _dt_pre.tzinfo is None:
            _dt_pre = _dt_pre.replace(tzinfo=timezone.utc)
        _pre_age = (datetime.now(timezone.utc) - _dt_pre).total_seconds()
    except Exception:
        _pre_age = 9999
    if _pre_age > 60:
        _live = _live_btc_price()
        if _live > 0:
            price = _live

    # UTC heartbeat -> local time + age
    _ts_raw = hb.get("ts", "")
    ts = "?"; data_age_s = 0
    try:
        _dt = datetime.fromisoformat(_ts_raw)
        if _dt.tzinfo is None:
            _dt = _dt.replace(tzinfo=timezone.utc)
        ts = _dt.astimezone().strftime("%H:%M:%S")
        data_age_s = int((datetime.now(timezone.utc) - _dt).total_seconds())
    except Exception:
        ts = _ts_raw[:19].replace("T", " ")

    if data_age_s <= 15:
        age_tag = f"  data {data_age_s}s old"
    elif data_age_s <= 60:
        age_tag = f"  ** data {data_age_s}s old **"
    else:
        _m, _s = data_age_s // 60, data_age_s % 60
        age_tag = f"  ** DATA {_m}m{_s}s OLD — CHECK CONNECTION **"

    C       = float(eng.get("C",           0))
    K       = float(eng.get("K",           0))
    ready   = bool(eng.get("ready",        False))
    samps   = int(eng.get("samples",       0))
    reg     = str(eng.get("regime",        "warmup"))
    cert    = float(eng.get("regime_cert", 0))
    agree   = float(eng.get("agreement",   0))
    volp    = float(eng.get("vol_penalty", 1))
    volf    = float(eng.get("vol_forecast",0))

    engine_live = bool(strat.get("coefficient_engine_enabled", False))
    coef_cfg    = strat.get("coefficient", {})
    tau_enter   = float(coef_cfg.get("tau_enter", 0.12))
    max_pos     = int(strat.get("max_open_positions", 7))
    version     = strat.get("version", "?")
    cb_state    = cb.get("state", "normal").upper()
    open_trades = hb.get("open_trades", [])

    balance, dolls = _calc_balance(closed)
    n       = len(closed)
    wins    = sum(1 for t in closed if _net(t) > 0)
    wr      = wins / n if n else 0.0

    total_unreal = 0.0
    for t in open_trades:
        ep  = float(t.get("entry_price", price))
        lev = float(t.get("leverage",    1.5))
        sz  = float(t.get("size",        0.15))
        raw = (price - ep) / ep if t.get("direction") == "long" else (ep - price) / ep
        total_unreal += raw * lev * sz * balance

    eq        = balance + total_unreal
    total_pnl = eq - STARTING
    eq_pnl_s  = f"+${total_pnl:,.2f}" if total_pnl >= 0 else f"-${abs(total_pnl):,.2f}"
    now_local = datetime.now().strftime("%H:%M:%S")
    stale_tag = f"  ** STALE since {_cache['stale_since']} **" if _cache["stale"] else ""

    D = "=" * TW

    # ── Clear + header ────────────────────────────────────────────────────────
    print("\033[2J\033[H", end="")
    print(D)
    mode = "LIVE" if engine_live else "SHADOW"
    print(f"  HERMES TRADING  [{mode}]  v{version}"
          f"{'':>{ TW - 36 - len(version) - len(mode)}}{now_local}")
    print(D)

    # ── Status bar ────────────────────────────────────────────────────────────
    print(_row(f"BTC  ${price:,.2f}   RSI {rsi:.1f}   {ts}{age_tag}{stale_tag}",
               f"Equity  ${eq:,.2f}  ({eq_pnl_s})"))
    if open_trades:
        pos_str = "   ".join(
            f"{t.get('direction','').upper()} @ ${float(t.get('entry_price',0)):,.0f}"
            for t in open_trades
        )
        print(_row(f"Status: IN TRADE  [{pos_str}]", f"Closed {n}   WR {wr*100:.0f}%"))
    else:
        print(_row("Status: waiting", f"Closed {n}   WR {wr*100:.0f}%"))
    print()

    # ── Engine ────────────────────────────────────────────────────────────────
    print(f"  {'─'*4} ENGINE {'─'*( TW - 12)}")
    if not ready:
        e_status = f"WARMING UP  ({samps}/200 samples,  ~{max(0,(200-samps)*INTERVAL)//60}min remaining)"
    elif engine_live:
        e_status = "LIVE AND READY"
    else:
        e_status = "SHADOW MODE"
    print(_row(f"Engine:  {e_status}", f"Regime: {reg.upper()}   cert={cert:.2f}   Circuit: {cb_state}"))

    direction = "LONG" if C > 0.05 else ("SHORT" if C < -0.05 else "FLAT")
    print(_row(f"C = {C:+.4f}  {bar(C, 24)}  {direction}",
               f"K = {K:.4f}   agree={agree:.2f}   vol_penalty={volp:.2f}   vol_fcast={volf:.5f}"))

    if ready:
        if engine_live:
            if abs(C) > tau_enter:
                signal_line = f"SIGNAL: ENTERING {('LONG' if C>0 else 'SHORT')}   |C|={abs(C):.4f} > tau={tau_enter}"
            else:
                signal_line = f"No entry   |C|={abs(C):.4f}  (need +{tau_enter-abs(C):.4f} to reach tau={tau_enter})"
        else:
            side = "LONG" if C > 0 else "SHORT"
            signal_line = (f"Shadow signal: {side}   |C|={abs(C):.4f}" if abs(C)>tau_enter
                           else f"Shadow: no signal   |C|={abs(C):.4f}")
    else:
        signal_line = f"Warming up — {200-samps} ticks to go"
    print(f"  {signal_line}")
    print()

    # ── Sub-signals ───────────────────────────────────────────────────────────
    print(f"  {'─'*4} SUB-SIGNALS  [{reg.upper()}] {'─'*(TW - 22 - len(reg))}")
    for key, label in [("mom","Momentum"),("rev","Mean-Rev"),
                       ("mic","Microstr"),("sen","Sentiment")]:
        v = float(sigs.get(key, 0))
        w = float(wts.get(key, 0))
        lean = "bullish" if v > 0.15 else ("bearish" if v < -0.15 else "neutral")
        print(f"  {label:10s}  {v:+.3f}  {bar(v, 20)}  w={w:.2f}  {lean}")
    print()

    # ── Market indicators ─────────────────────────────────────────────────────
    print(f"  {'─'*4} MARKET {'─'*(TW - 13)}")
    ob_lbl = "buy pressure" if ob_imb>0.1 else ("sell pressure" if ob_imb<-0.1 else "balanced")
    bb_lbl = "overbought"   if bb>0.8      else ("oversold"      if bb<0.2      else f"{bb:.2f}")
    print(_row(f"RSI {rsi:5.1f}   MACD {macd:+.4f}   ATR ${atr:.2f}   BB {bb_lbl}",
               f"OB imbalance {ob_imb:+.3f}  ({ob_lbl})"))
    print()

    # ── Open positions ────────────────────────────────────────────────────────
    print(f"  {'─'*4} OPEN POSITIONS  ({len(open_trades)}/{max_pos}) {'─'*(TW - 28)}")
    if open_trades:
        hdr = f"  {'DIR':<6} {'ENTRY':>12}  {'LEV':>5}  {'SIZE':>5}  {'UNREALISED %':>13}  {'UNREALISED $':>13}  {'C':>7}  {'K':>6}"
        print(hdr)
        print("  " + "-" * (TW - 2))
        for t in open_trades:
            ep      = float(t.get("entry_price", price))
            lev     = float(t.get("leverage",    1.5))
            sz      = float(t.get("size",        0.15))
            ec      = t.get("entry_C"); ek = t.get("entry_K")
            raw     = (price-ep)/ep if t.get("direction")=="long" else (ep-price)/ep
            unr_pct = raw * lev * 100
            unr_usd = raw * lev * sz * balance
            c_str   = f"{ec:+.3f}" if ec is not None else "  -  "
            k_str   = f"{ek:.3f}"  if ek is not None else "  - "
            sgn     = "+" if unr_pct >= 0 else ""
            print(f"  {t.get('direction','').upper():<6} ${ep:>11,.2f}  {lev:>4.1f}x  {sz*100:>4.0f}%"
                  f"  {sgn}{unr_pct:>12.2f}%  {sgn}${abs(unr_usd):>11,.0f}  {c_str:>7}  {k_str:>6}")
        sgn = "+" if total_unreal >= 0 else ""
        print(f"  {'':>56} Total  {sgn}${abs(total_unreal):>11,.0f}")
    else:
        print("  No open positions")
    print()

    # ── Performance summary ───────────────────────────────────────────────────
    print(f"  {'─'*4} PERFORMANCE {'─'*(TW - 18)}")
    if n > 0:
        win_pnls  = [_net(t) for t in closed if _net(t) > 0]
        loss_pnls = [_net(t) for t in closed if _net(t) <= 0]
        avg_win   = sum(win_pnls)  / len(win_pnls)  if win_pnls  else 0.0
        avg_loss  = sum(loss_pnls) / len(loss_pnls) if loss_pnls else 0.0
        best_usd  = max(dolls) if dolls else 0.0
        worst_usd = min(dolls) if dolls else 0.0
        cum = 1.0; pk = 1.0; dd = 0.0
        for t in closed:
            cum *= (1 + _net(t)); pk = max(pk, cum); dd = max(dd, (pk-cum)/pk)

        pnl_str = f"+${total_pnl:,.2f}" if total_pnl >= 0 else f"-${abs(total_pnl):,.2f}"
        unr_str = f"+${total_unreal:,.0f}" if total_unreal >= 0 else f"-${abs(total_unreal):,.0f}"
        print(_row(f"Starting: ${STARTING:,.0f}   Realised: ${balance:,.2f}   "
                   f"Unrealised: {unr_str}",
                   f"Total equity: ${eq:,.2f}  ({pnl_str})"))
        print(_row(f"Trades: {n}   Wins: {wins}   Losses: {n-wins}   Win rate: {wr*100:.0f}%",
                   f"Avg win: {avg_win*100:+.2f}%   Avg loss: {avg_loss*100:+.2f}%   Max DD: -{dd*100:.2f}%"))
        print(_row(f"Best trade:  +${best_usd:,.0f}",
                   f"Worst trade:  -${abs(worst_usd):,.0f}"))
    else:
        print(f"  No closed trades yet   Starting balance: ${STARTING:,.0f}")
    print()

    # ── Trades table ─────────────────────────────────────────────────────────
    print(f"  {'─'*4} ALL TRADES  (newest first) {'─'*(TW - 32)}")
    if not closed:
        print("  No closed trades yet.")
    else:
        # Table header
        TH = (f"  {'#':>4}  {'DIR':<6}  {'ENTRY':>12}  {'EXIT':>12}  "
              f"{'PNL %':>8}  {'PNL $':>10}  {'BALANCE':>12}  "
              f"{'REASON':<18}  {'LEVERAGE':>8}  {'CLOSED'}")
        print(TH)
        print("  " + "-" * (TW - 2))
        running = STARTING
        running_list = []
        for d in dolls:
            running += d
            running_list.append(running)
        for idx in range(n - 1, -1, -1):
            t    = closed[idx]
            d    = dolls[idx]
            bal  = running_list[idx]
            pct    = _net(t)
            pct_s  = f"+{pct*100:.2f}%" if pct >= 0 else f"-{abs(pct*100):.2f}%"
            d_s    = f"+${d:,.0f}"      if d >= 0    else f"-${abs(d):,.0f}"
            exit_t = t.get("exit_time", "")[:16].replace("T", " ")
            reason = t.get("exit_reason", "").replace("_", " ")[:18]
            lev    = t.get("leverage", 1)
            print(f"  {idx+1:>4}  {t.get('direction','').upper():<6}  "
                  f"${t.get('entry_price',0):>11,.2f}  ${t.get('exit_price',0):>11,.2f}  "
                  f"{pct_s:>9}  {d_s:>11}  "
                  f"${bal:>11,.2f}  {reason:<18}  {lev:>6}x    {exit_t}")
    print()
    print(f"  Refreshing every {INTERVAL}s   Ctrl+C to stop")
    print("=" * TW)


# ── CSV writers ───────────────────────────────────────────────────────────────

def write_csvs(state: dict, closed: list) -> None:
    hb    = state.get("heartbeat", {})
    strat = state.get("strategy", {})
    goal  = state.get("goal",     {})
    model = state.get("model",    {})
    eng   = hb.get("engine", {})

    price         = float(hb.get("price", 0))
    rsi           = float(hb.get("rsi",   0))
    open_trades   = hb.get("open_trades", [])
    starting      = float(goal.get("starting_balance", STARTING))
    ecfg          = strat.get("entry", {})
    sl            = float(strat.get("stop_loss_pct", 0.5))
    tp            = float(strat.get("take_profit_pct", 0.0))
    ver           = hb.get("strategy_version", "?")
    ts            = hb.get("ts", "")
    fails         = hb.get("consecutive_failures", 0)

    balance, dolls = _calc_balance(closed)
    net_pnl        = balance - starting

    # Unrealised
    unrealised = 0.0
    open_details = []
    for t in open_trades:
        ep  = float(t.get("entry_price", price) or price)
        lev = float(t.get("leverage", 1.0))
        sz  = float(t.get("size") or t.get("position_size_r", 0.15))
        if ep <= 0 or price <= 0:
            continue
        raw   = (price-ep)/ep if t.get("direction")=="long" else (ep-price)/ep
        gross = raw * lev
        dol   = gross * sz * balance
        unrealised += dol
        open_details.append({
            "id": t.get("id","?"), "direction": t.get("direction","long").upper(),
            "entry": ep, "raw_pct": raw, "gross_pct": gross,
            "dollar": dol, "leverage": lev, "size": sz,
        })

    eq              = balance + unrealised
    total_pnl       = eq - starting

    # ── position.csv ─────────────────────────────────────────────────────────
    rows = [
        ["Field",                "Value"],
        ["Updated",              datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
        ["Source",               RAILWAY_URL],
        ["Asset",                hb.get("asset", "BTC/USDT")],
        ["Current price",        f"${price:,.2f}"],
        ["RSI",                  round(rsi, 2)],
        ["MACD hist",            round(hb.get("macd_hist", 0), 4)],
        ["BB %B",                round(hb.get("bb_pct", 0), 4)],
        ["OB imbalance",         round(hb.get("ob_imbalance", 0), 4)],
        ["Engine C",             round(float(eng.get("C", 0)), 4)],
        ["Engine K",             round(float(eng.get("K", 0)), 4)],
        ["Regime",               eng.get("regime", "?")],
        [""],
        ["── OPEN POSITIONS ──", f"{len(open_details)} open"],
    ]
    if open_details:
        for i, d in enumerate(open_details, 1):
            rows += [
                [f"  Position {i}",         f"{d['direction']}  {d['size']*100:.0f}%  {d['leverage']:.1f}x lev"],
                [f"  Entry price",           f"${d['entry']:,.2f}"],
                [f"  Move vs entry",         f"{'+' if d['raw_pct']>=0 else ''}{d['raw_pct']*100:.3f}%"],
                [f"  Gross PnL (levered)",   f"{'+' if d['gross_pct']>=0 else ''}{d['gross_pct']*100:.3f}%"],
                [f"  Dollar PnL",            _usd(d["dollar"])],
                [f"  Stop loss at",
                 f"${d['entry']*(1+sl/100) if d['direction']=='SHORT' else d['entry']*(1-sl/100):,.2f}"],
            ]
    else:
        rows.append(["  No open positions", ""])
    rows += [
        [""],
        ["── BALANCE ──",        ""],
        ["Starting balance",     f"${starting:,.2f}"],
        ["Realised PnL",         _usd(net_pnl)],
        ["Unrealised PnL",       _usd(unrealised)],
        ["Total PnL $",          _usd(total_pnl)],
        ["Total PnL %",          _pct(total_pnl / starting)],
        ["Effective balance",    f"${eq:,.2f}"],
        [""],
        ["── ENGINE ──",         ""],
        ["C (direction)",        round(float(eng.get("C", 0)), 4)],
        ["K (conviction)",       round(float(eng.get("K", 0)), 4)],
        ["Regime",               eng.get("regime", "?")],
        ["Regime certainty",     round(float(eng.get("regime_cert", 0)), 4)],
        ["Engine enabled",       strat.get("coefficient_engine_enabled", False)],
        ["Samples",              eng.get("samples", 0)],
        [""],
        ["── STRATEGY ──",       ""],
        ["Long entry RSI <",     ecfg.get("long_threshold", 55)],
        ["Short entry RSI >",    ecfg.get("short_threshold", 65)],
        ["Max positions",        strat.get("max_open_positions", 7)],
        ["Stop loss",            f"-{sl}%"],
        ["Take profit",          f"+{tp}%"],
        ["Strategy version",     f"v{ver}"],
        ["Last tick",            _age(hb.get("ts",""))],
        ["Consecutive failures", fails],
        ["Agent",                "Healthy" if fails == 0 else "Degraded"],
    ]
    _write_csv("position.csv", rows)

    # ── stats.csv ─────────────────────────────────────────────────────────────
    if not closed:
        stats = [
            ["Field",             "Value"],
            ["Updated",           datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
            ["Starting balance",  f"${starting:,.2f}"],
            ["Effective balance", f"${eq:,.2f}"],
            ["Realised PnL",      _usd(net_pnl)],
            ["Unrealised PnL",    _usd(unrealised)],
            ["Total PnL $",       _usd(total_pnl)],
            ["Total PnL %",       _pct(total_pnl / starting)],
            ["Trades closed",     0],
            ["Open positions",    len(open_details)],
        ]
    else:
        wins    = [t for t in closed if _net(t) > 0]
        losses  = [t for t in closed if _net(t) <= 0]
        longs   = [t for t in closed if t.get("direction") == "long"]
        shorts  = [t for t in closed if t.get("direction") == "short"]
        wr      = len(wins) / len(closed) * 100
        best_i  = max(range(len(closed)), key=lambda i: _net(closed[i]))
        worst_i = min(range(len(closed)), key=lambda i: _net(closed[i]))

        cum = 1.0; pk = 1.0; dd = 0.0
        for t in closed:
            cum *= (1 + _net(t))
            pk   = max(pk, cum)
            dd   = max(dd, (pk - cum) / pk)

        score_counts: dict = {}
        for t in closed:
            sc = t.get("entry_score", "eng")
            score_counts[sc] = score_counts.get(sc, 0) + 1
        score_str = "  ".join(f"{k}×{v}" for k, v in sorted(score_counts.items(), key=lambda x: str(x[0])))

        stats = [
            ["Field",             "Value"],
            ["Updated",           datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
            ["Starting balance",  f"${starting:,.2f}"],
            ["Realised balance",  f"${balance:,.2f}"],
            ["Unrealised PnL",    _usd(unrealised)],
            ["Total PnL $",       _usd(total_pnl)],
            ["Total PnL %",       _pct(total_pnl / starting)],
            ["Effective balance", f"${eq:,.2f}"],
            ["Open positions",    len(open_details)],
            ["Trades closed",     f"{len(closed)}  ({len(wins)}W / {len(losses)}L)"],
            ["Long / Short",      f"{len(longs)}L / {len(shorts)}S"],
            ["Win rate",          f"{wr:.0f}%"],
            ["Score dist",        score_str or "n/a"],
            ["Best trade $",      _usd(dolls[best_i])],
            ["Worst trade $",     _usd(dolls[worst_i])],
            ["Max drawdown",      f"-{dd * 100:.2f}%"],
        ]
    _write_csv("stats.csv", stats)

    # ── trades.csv ────────────────────────────────────────────────────────────
    header = ["Trade ID","Direction","Score","Entry Price","Exit Price",
              "PnL $","PnL %","Balance After","Exit Reason",
              "Strategy","Leverage","Entry Time","Exit Time"]
    if not closed:
        trade_rows = [header, ["No closed trades yet"] + [""]*12]
    else:
        running = starting
        rows2 = []
        for i, t in enumerate(closed):
            running += dolls[i]
            rows2.append([
                t.get("id",""),
                t.get("direction","").upper(),
                f"{t.get('entry_score','eng')}",
                f"${t.get('entry_price',0):,.2f}",
                f"${t.get('exit_price',0):,.2f}",
                _usd(dolls[i]),
                _pct(_net(t)),
                f"${running:,.2f}",
                t.get("exit_reason","").replace("_"," "),
                "v" + str(t.get("strategy_version","?")),
                f'{t.get("leverage",1)}x',
                t.get("entry_time",""),
                t.get("exit_time",""),
            ])
        trade_rows = [header] + list(reversed(rows2))
    _write_csv("trades.csv", trade_rows)

    # ── model.csv ─────────────────────────────────────────────────────────────
    if model.get("status") == "ok":
        mc  = model.get("monte_carlo", {})
        k   = model.get("kelly", {})
        reg = model.get("regime", {})
        st  = model.get("stats", {})
        rec = model.get("recommendations", [])
        model_rows = [
            ["Field",                        "Value"],
            ["Updated",                      datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
            ["Signal",                       model.get("signal","").replace("_"," ").upper()],
            ["Action",                       model.get("action","").upper()],
            ["Confidence",                   f"{model.get('confidence',0):.1%}"],
            [""],
            ["── REGIME ──",                 ""],
            ["Market regime",                reg.get("regime","")],
            ["Regime quality",               f"{reg.get('quality',0):.1%}"],
            ["ATR %",                        f"{reg.get('atr_pct',0):.3f}%"],
            [""],
            ["── MONTE CARLO ──",            ""],
            ["Paths",                        f"{mc.get('n_paths',0):,}"],
            ["P5  (bad case)",               f"${mc.get('equity_p5',0):,.0f}"],
            ["P50 (median)",                 f"${mc.get('equity_p50',0):,.0f}"],
            ["P95 (good case)",              f"${mc.get('equity_p95',0):,.0f}"],
            ["Expected return",              _pct(mc.get("expected_return",0))],
            ["Max drawdown P95",             f"{mc.get('max_dd_p95',0)*100:.2f}%"],
            ["Risk of ruin",                 f"{mc.get('ruin_probability',0):.1%}"],
            [""],
            ["── KELLY ──",                  ""],
            ["Win probability",              f"{k.get('win_probability',0):.1%}"],
            ["Payoff ratio",                 f"{k.get('payoff_ratio',0):.2f}x"],
            ["Kelly quarter (recommended)",  f"{k.get('kelly_quarter',0):.1%}"],
            [""],
            ["── TRADE STATS ──",            ""],
            ["Win rate",                     f"{st.get('win_rate',0):.1%}"],
            ["Avg win",                      f"{st.get('avg_win_pct',0):+.3f}%"],
            ["Avg loss",                     f"{st.get('avg_loss_pct',0):+.3f}%"],
            ["Sharpe",                       f"{st.get('sharpe',0):.3f}"],
            ["Max drawdown (actual)",        f"{st.get('max_drawdown',0)*100:.2f}%"],
            [""],
            ["── RECOMMENDATIONS ──",        ""],
        ]
        if not rec:
            model_rows.append(["No recommendations", "Within normal bounds"])
        else:
            for r in rec:
                model_rows.append([
                    f"[{r.get('urgency','').upper()}] {r.get('field','')}",
                    f"{r.get('current','')} → {r.get('suggested','')}",
                ])
                model_rows.append(["  Reason", r.get("reason","")])
    else:
        model_rows = [
            ["Field",  "Value"],
            ["Status", model.get("status","model not yet run")],
            ["Note",   "Need 3+ closed trades to activate"],
        ]
    _write_csv("model.csv", model_rows)


# ── Excel spreadsheet ─────────────────────────────────────────────────────────

def _xl_font(bold=False, size=10, color="000000") -> Font:
    return Font(name="Arial", bold=bold, size=size, color=color)

def _xl_fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)

def _xl_border_bottom() -> Border:
    thin = Side(style="thin", color="CCCCCC")
    return Border(bottom=thin)

def _xl_header_border() -> Border:
    thin  = Side(style="thin", color="999999")
    thick = Side(style="medium", color="555555")
    return Border(bottom=thick, top=thin)

def _xl_set_col_width(ws, col: int, width: float) -> None:
    ws.column_dimensions[get_column_letter(col)].width = width


def write_xlsx(state: dict, closed: list) -> None:
    """Write hermes_trades.xlsx with Trades and Performance sheets."""
    hb      = state.get("heartbeat", {})
    strat   = state.get("strategy",  {})
    goal    = state.get("goal",      {})
    eng     = hb.get("engine", {})

    price         = float(hb.get("price", 0))
    open_trades   = hb.get("open_trades", [])
    starting      = float(goal.get("starting_balance", STARTING))
    engine_live   = bool(strat.get("coefficient_engine_enabled", False))
    max_pos       = int(strat.get("max_open_positions", 7))

    balance, dolls = _calc_balance(closed)
    net_pnl        = balance - starting
    n              = len(closed)
    wins           = sum(1 for t in closed if _net(t) > 0)
    wr             = wins / n if n else 0.0

    # Unrealised across open positions
    unrealised = 0.0
    for t in open_trades:
        ep  = float(t.get("entry_price", price) or price)
        lev = float(t.get("leverage", 1.5))
        sz  = float(t.get("size") or t.get("position_size_r", 0.15))
        if ep > 0 and price > 0:
            raw = (price - ep) / ep if t.get("direction") == "long" else (ep - price) / ep
            unrealised += raw * lev * sz * balance
    eq = balance + unrealised

    # Timestamp freshness
    _ts_raw = hb.get("ts", "")
    try:
        _dt = datetime.fromisoformat(_ts_raw)
        if _dt.tzinfo is None:
            _dt = _dt.replace(tzinfo=timezone.utc)
        hb_local = _dt.astimezone().strftime("%Y-%m-%d %H:%M:%S")
        data_age = int((datetime.now(timezone.utc) - _dt).total_seconds())
        age_str  = f"{data_age}s old"
    except Exception:
        hb_local = _ts_raw[:19].replace("T", " ")
        age_str  = "?"

    updated = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    wb = Workbook()

    # ── Sheet 1: Trades ───────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Trades"
    ws.sheet_view.showGridLines = False

    # Freeze top 2 rows (title + header)
    ws.freeze_panes = "A3"

    # Title row
    ws.merge_cells("A1:M1")
    title_cell = ws["A1"]
    title_cell.value = f"Hermes Trades   |   Updated {updated}   |   Data {age_str}"
    title_cell.font  = _xl_font(bold=True, size=11)
    title_cell.fill  = _xl_fill("FFFFFF")
    title_cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 22

    # Column headers
    HEADERS = [
        "#", "Direction", "Score", "Entry Price", "Exit Price",
        "PnL $", "PnL %", "Balance After", "Exit Reason",
        "Strategy", "Leverage", "Entry Time", "Exit Time",
    ]
    COL_WIDTHS = [5, 11, 8, 14, 14, 14, 10, 16, 18, 11, 10, 20, 20]

    for col, (h, w) in enumerate(zip(HEADERS, COL_WIDTHS), 1):
        c = ws.cell(row=2, column=col, value=h)
        c.font      = _xl_font(bold=True, size=10)
        c.fill      = _xl_fill("F2F2F2")
        c.border    = _xl_header_border()
        c.alignment = Alignment(horizontal="center", vertical="center")
        _xl_set_col_width(ws, col, w)
    ws.row_dimensions[2].height = 18

    # Data rows (newest first)
    if not closed:
        ws.merge_cells("A3:M3")
        c = ws.cell(row=3, column=1, value="No closed trades yet.")
        c.font = _xl_font(size=10, color="888888")
        c.alignment = Alignment(horizontal="left")
    else:
        running = starting
        for row_i, (t, d) in enumerate(zip(closed, dolls), start=0):
            running_before = running
            running += d
            # Reverse: newest at top
            display_row = len(closed) - row_i
            data_row    = 2 + display_row  # row 3 = trade #N (newest), row N+2 = trade #1
        # Write reversed
        running = starting
        running_list = []
        for d in dolls:
            running += d
            running_list.append(running)

        for idx in range(n - 1, -1, -1):
            t   = closed[idx]
            d   = dolls[idx]
            bal = running_list[idx]
            xr  = 2 + (n - 1 - idx) + 1   # row index in sheet (row 3 = newest)

            pnl_pct = _net(t)
            is_win  = pnl_pct > 0
            row_fill = _xl_fill("F6FBF6") if is_win else _xl_fill("FBF6F6")
            pnl_color = "1A7A1A" if is_win else "B02020"

            values = [
                idx + 1,
                t.get("direction", "").upper(),
                t.get("entry_score", "eng"),
                t.get("entry_price", 0),
                t.get("exit_price", 0),
                d,
                pnl_pct,
                bal,
                t.get("exit_reason", "").replace("_", " "),
                "v" + str(t.get("strategy_version", "?")),
                f'{t.get("leverage", 1)}x',
                t.get("entry_time", ""),
                t.get("exit_time", ""),
            ]
            # Alignment per column: center most, left for text
            aligns = ["center","center","center","right","right",
                      "right","right","right","left","center","center","center","center"]
            for col, (val, al) in enumerate(zip(values, aligns), 1):
                c = ws.cell(row=xr, column=col, value=val)
                c.fill      = row_fill
                c.border    = _xl_border_bottom()
                c.alignment = Alignment(horizontal=al, vertical="center")
                # Formatting
                if col == 4 or col == 5:   # prices
                    c.number_format = '"$"#,##0.00'
                    c.font = _xl_font(size=10)
                elif col == 6:             # PnL $
                    c.number_format = '"$"#,##0.00;"-$"#,##0.00'
                    c.font = _xl_font(bold=True, size=10, color=pnl_color)
                elif col == 7:             # PnL %
                    c.number_format = '+0.00%;-0.00%'
                    c.font = _xl_font(bold=True, size=10, color=pnl_color)
                elif col == 8:             # balance after
                    c.number_format = '"$"#,##0.00'
                    c.font = _xl_font(size=10)
                elif col == 2:             # direction
                    dir_color = "1A4FA0" if val == "LONG" else "A01A1A"
                    c.font = _xl_font(bold=True, size=10, color=dir_color)
                else:
                    c.font = _xl_font(size=10)
            ws.row_dimensions[xr].height = 16

    # ── Sheet 2: Performance ─────────────────────────────────────────────────
    ws2 = wb.create_sheet("Performance")
    ws2.sheet_view.showGridLines = False

    def perf_title(row, text):
        ws2.merge_cells(f"A{row}:B{row}")
        c = ws2.cell(row=row, column=1, value=text)
        c.font      = _xl_font(bold=True, size=10)
        c.fill      = _xl_fill("F2F2F2")
        c.border    = _xl_header_border()
        c.alignment = Alignment(horizontal="left", vertical="center")
        ws2.row_dimensions[row].height = 18

    def perf_row(row, label, value, fmt=None, color=None):
        lc = ws2.cell(row=row, column=1, value=label)
        lc.font      = _xl_font(size=10)
        lc.fill      = _xl_fill("FFFFFF")
        lc.border    = _xl_border_bottom()
        lc.alignment = Alignment(horizontal="left", vertical="center")

        vc = ws2.cell(row=row, column=2, value=value)
        vc.font      = _xl_font(bold=False, size=10, color=color or "000000")
        vc.fill      = _xl_fill("FFFFFF")
        vc.border    = _xl_border_bottom()
        vc.alignment = Alignment(horizontal="right", vertical="center")
        if fmt:
            vc.number_format = fmt
        ws2.row_dimensions[row].height = 16

    ws2.column_dimensions["A"].width = 28
    ws2.column_dimensions["B"].width = 20

    # Title
    ws2.merge_cells("A1:B1")
    t1 = ws2["A1"]
    t1.value     = f"Performance   |   Updated {updated}"
    t1.font      = _xl_font(bold=True, size=11)
    t1.fill      = _xl_fill("FFFFFF")
    t1.alignment = Alignment(horizontal="left", vertical="center")
    ws2.row_dimensions[1].height = 22

    r = 2

    # Live feed
    perf_title(r, "LIVE"); r += 1
    perf_row(r, "BTC Price",        price,      '"$"#,##0.00'); r += 1
    perf_row(r, "Last heartbeat",   hb_local                 ); r += 1
    perf_row(r, "Data age",         age_str                  ); r += 1
    perf_row(r, "Engine mode",      "LIVE" if engine_live else "SHADOW"); r += 1
    perf_row(r, "Engine C",         float(eng.get("C", 0)),  '+0.0000'); r += 1
    perf_row(r, "Engine K",         float(eng.get("K", 0)),  '0.0000' ); r += 1
    perf_row(r, "Regime",           eng.get("regime", "?")             ); r += 1
    perf_row(r, "Open positions",   f"{len(open_trades)}/{max_pos}"    ); r += 1
    r += 1

    # Balance
    pnl_color_val = "1A7A1A" if eq >= starting else "B02020"
    perf_title(r, "BALANCE"); r += 1
    perf_row(r, "Starting balance", starting,   '"$"#,##0.00'); r += 1
    perf_row(r, "Realised balance", balance,    '"$"#,##0.00'); r += 1
    perf_row(r, "Unrealised PnL",   unrealised, '"$"+#,##0.00;"-$"#,##0.00'); r += 1
    perf_row(r, "Total equity",     eq,         '"$"#,##0.00'); r += 1
    pnl_cell = ws2.cell(row=r, column=2)
    perf_row(r, "Total PnL $",      eq - starting, '"$"+#,##0.00;"-$"#,##0.00',
             color=pnl_color_val); r += 1
    perf_row(r, "Total PnL %",      (eq - starting) / starting if starting else 0,
             '+0.00%;-0.00%', color=pnl_color_val); r += 1
    r += 1

    # Stats
    perf_title(r, "TRADE STATISTICS"); r += 1
    perf_row(r, "Closed trades",    n); r += 1
    perf_row(r, "Wins",             wins); r += 1
    perf_row(r, "Losses",           n - wins); r += 1
    perf_row(r, "Win rate",         wr, '0.0%'); r += 1

    if n > 0:
        win_pnls  = [_net(t) for t in closed if _net(t) > 0]
        loss_pnls = [_net(t) for t in closed if _net(t) <= 0]
        avg_win   = sum(win_pnls)  / len(win_pnls)  if win_pnls  else 0.0
        avg_loss  = sum(loss_pnls) / len(loss_pnls) if loss_pnls else 0.0
        best_pct  = max(_net(t) for t in closed)
        worst_pct = min(_net(t) for t in closed)
        best_usd  = max(dolls)
        worst_usd = min(dolls)

        cum = 1.0; pk = 1.0; dd = 0.0
        for t in closed:
            cum *= (1 + _net(t))
            pk   = max(pk, cum)
            dd   = max(dd, (pk - cum) / pk)

        perf_row(r, "Avg win %",     avg_win,   '+0.00%;-0.00%', color="1A7A1A"); r += 1
        perf_row(r, "Avg loss %",    avg_loss,  '+0.00%;-0.00%', color="B02020"); r += 1
        perf_row(r, "Best trade $",  best_usd,  '"$"+#,##0.00;"-$"#,##0.00',  color="1A7A1A"); r += 1
        perf_row(r, "Worst trade $", worst_usd, '"$"+#,##0.00;"-$"#,##0.00',  color="B02020"); r += 1
        perf_row(r, "Max drawdown",  -dd,       '0.00%',                       color="B02020"); r += 1

    path = BASE / "hermes_trades.xlsx"
    try:
        wb.save(path)
    except PermissionError:
        # File is open in Excel — skip silently this tick
        pass


# ── Main ──────────────────────────────────────────────────────────────────────

ERROR_LOG = BASE / "hermes_error.log"


def main() -> None:
    print(f"HERMES Monitor — connecting to {RAILWAY_URL}")
    print("Writing position.csv / stats.csv / trades.csv / model.csv / hermes_trades.xlsx\n")
    while True:
        try:
            state, closed = fetch_all()
            render_terminal(state, closed)
            write_csvs(state, closed)
            write_xlsx(state, closed)
        except KeyboardInterrupt:
            print("\nStopped.")
            break
        except Exception as exc:
            now = datetime.now().strftime("%H:%M:%S")
            tb_str = traceback.format_exc()
            # Write to log file (persists across screen clears)
            with open(ERROR_LOG, "a") as f:
                f.write(f"\n[{now}] {exc}\n{tb_str}\n")
            # Print WITHOUT triggering a screen clear
            print(f"\n  [{now}] error — {exc}")
            print(tb_str)
            print(f"  (full traceback saved to hermes_error.log)")
        time.sleep(INTERVAL)


if __name__ == "__main__":
    main()
